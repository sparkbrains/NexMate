from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from apps.db import get_connection, utc_now


# The prompt pack lives in data/prompt_pack.xlsx (columns: id, category,
# text) rather than in code, so it can be edited/extended without a
# deploy. Stable `id` per row so answers stay correctly attributed even
# if rows are reordered or edited later -- rotation just cycles through
# whatever's currently in the sheet. Regenerate/extend it with
# scripts/generate_prompt_pack_xlsx.py.
_PROMPT_PACK_XLSX = Path(__file__).resolve().parents[3] / "data" / "prompt_pack.xlsx"

_prompt_pack_cache: list[dict[str, str]] | None = None


def _load_prompt_pack() -> list[dict[str, str]]:
    global _prompt_pack_cache
    if _prompt_pack_cache is not None:
        return _prompt_pack_cache

    if not _PROMPT_PACK_XLSX.exists():
        raise RuntimeError(f"Prompt pack file not found: {_PROMPT_PACK_XLSX}")

    wb = load_workbook(_PROMPT_PACK_XLSX, read_only=True, data_only=True)
    try:
        ws = wb["Prompts"] if "Prompts" in wb.sheetnames else wb.active
        pack = [
            {"id": str(row[0]).strip(), "category": str(row[1]).strip(), "text": str(row[2]).strip()}
            for row in ws.iter_rows(min_row=2, values_only=True)
            if row and row[0] and row[1] and row[2]
        ]
    finally:
        wb.close()

    if not pack:
        raise RuntimeError(f"No prompts found in {_PROMPT_PACK_XLSX}")

    _prompt_pack_cache = pack
    return pack


def get_prompt_for_date(d: date) -> dict[str, str]:
    """Deterministic day -> prompt mapping. Same prompt for everyone on a
    given calendar day, cycles through the whole pack, wraps around."""
    pack = _load_prompt_pack()
    idx = d.toordinal() % len(pack)
    return pack[idx]


def get_todays_prompt(user_id: int) -> dict[str, Any]:
    today = date.today()
    p = get_prompt_for_date(today)

    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT prompt_id, answer_text
                FROM prompt_pack_answers
                WHERE user_id = %s AND answered_date = %s
                """,
                (user_id, today),
            )
            row = cur.fetchone()

    return {
        "prompt_id": p["id"],
        "prompt_text": p["text"],
        "category": p["category"],
        "answered": row is not None,
        "answer_text": row["answer_text"] if row else None,
    }


async def save_prompt_answer(user_id: int, prompt_id: str, answer_text: str) -> dict[str, Any]:
    today = date.today()
    p = get_prompt_for_date(today)

    if p["id"] != prompt_id:
        # Guards against a stale client submitting yesterday's prompt_id
        # after midnight, or a tampered request.
        raise ValueError("This isn't today's prompt anymore")

    cleaned = answer_text.strip()
    if not cleaned:
        raise ValueError("Answer can't be empty")
    if len(cleaned) > 5000:
        raise ValueError("Answer is too long")

    now = utc_now()
    # Insert or update the answer in the DB
    with get_connection(autocommit=True) as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO prompt_pack_answers
                    (user_id, prompt_id, prompt_text, category, answer_text, answered_date, created_at)
                VALUES (%s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (user_id, answered_date) DO UPDATE SET
                    prompt_id = EXCLUDED.prompt_id,
                    prompt_text = EXCLUDED.prompt_text,
                    category = EXCLUDED.category,
                    answer_text = EXCLUDED.answer_text,
                    created_at = EXCLUDED.created_at
                RETURNING id, created_at
                """,
                (user_id, prompt_id, p["text"], p["category"], cleaned, today, now),
            )
            row = cur.fetchone()

    # Trigger profile summary generation in background
    import asyncio
    from apps.api.services.user_profile_service import get_or_refresh_user_profile
    asyncio.create_task(get_or_refresh_user_profile(user_id))

    return {
        "id": int(row["id"]),
        "prompt_id": prompt_id,
        "prompt_text": p["text"],
        "answer_text": cleaned,
        "created_at": row["created_at"].isoformat(),
    }

def list_all_prompts(user_id: int) -> dict[str, Any]:
    """Every prompt in the pack, grouped by category, with this user's
    answer (if any) merged in -- for the Prompt Packs tab, which shows
    the whole pack rather than just today's single prompt."""
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT prompt_id, answer_text, answered_date
                FROM prompt_pack_answers
                WHERE user_id = %s
                ORDER BY answered_date DESC
                """,
                (user_id,),
            )
            rows = cur.fetchall()

    # Most recent answer wins if a prompt was ever answered more than once
    # (rotation wraps after 15 days, so this is possible over time).
    answered_by_prompt: dict[str, dict[str, Any]] = {}
    for r in rows:
        if r["prompt_id"] not in answered_by_prompt:
            answered_by_prompt[r["prompt_id"]] = {
                "answer_text": r["answer_text"],
                "answered_date": r["answered_date"].isoformat(),
            }

    pack = _load_prompt_pack()
    categories: dict[str, list[dict[str, Any]]] = {}
    for p in pack:
        ans = answered_by_prompt.get(p["id"])
        entry = {
            "prompt_id": p["id"],
            "prompt_text": p["text"],
            "category": p["category"],
            "answered": ans is not None,
            "answer_text": ans["answer_text"] if ans else None,
            "answered_date": ans["answered_date"] if ans else None,
        }
        categories.setdefault(p["category"], []).append(entry)

    return {
        "categories": [
            {"category": cat, "prompts": prompts}
            for cat, prompts in categories.items()
        ],
        "total": len(pack),
        "answered_count": len(answered_by_prompt),
    }


def list_prompt_answers(user_id: int, limit: int = 90) -> list[dict[str, Any]]:
    """Recent answers, most recent first -- for a future 'your answers over
    time' view, and eventually the profile-compression step."""
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT id, prompt_id, prompt_text, category, answer_text, answered_date, created_at
                FROM prompt_pack_answers
                WHERE user_id = %s
                ORDER BY answered_date DESC
                LIMIT %s
                """,
                (user_id, limit),
            )
            rows = cur.fetchall()

    return [
        {
            "id": int(r["id"]),
            "prompt_id": r["prompt_id"],
            "prompt_text": r["prompt_text"],
            "category": r["category"],
            "answer_text": r["answer_text"],
            "answered_date": r["answered_date"].isoformat(),
            "created_at": r["created_at"].isoformat(),
        }
        for r in rows
    ]